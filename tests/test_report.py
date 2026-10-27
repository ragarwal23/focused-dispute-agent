"""The release workbook: sheets and titles, a null draft, an exception row, a scenario with no run, a non-report input."""

import json

import openpyxl
import pytest

from dispute_agent.evaluation.golden import build_cases
from dispute_agent.evaluation.report import NO_DRAFT, load_eval, render, write_report
from dispute_agent.__main__ import main


def normal_row(s):
    return {
        "scenario_id": s.scenario_id, "split": s.split, "trap": s.trap, "wording": s.wording, "note_driven": s.note_driven,
        "category_pred": s.expected_category, "account_attempts": 1, "policy_attempts": 1,
        "case": {"case_id": "C-004", "created": True, "route": s.expected_route},
        "checks": {"tool_calls": 3, "route_final": s.expected_route},
        "findings": {"relevant_event_ids": ["E-0081", "E-0117", "E-0120"], "disputed_event_id": "E-0117",
                     "candidate_event_ids": [], "summary": "cancelled autopay, fee stands"},
        "derived": {"payment_state": "posted", "refund_state": "not_applicable"},
        "sync": {"issues": []},
        "draft": {"review_route": s.expected_route, "event_ids": ["E-0081", "E-0117"], "policy_source_ids": ["AP.v1.s4", "LF.v3.s4"],
                  "confidence": "high", "open_questions": [], "case_note": "The fee stands."},
        "gate": {"route": s.expected_route, "gate_rejected": False, "gate_reasons": []},
        "tool_errors": [],
    }


def null_draft_row(s):
    row = normal_row(s)
    row.update(draft=None, case={"case_id": "C-005", "created": True, "route": "manual_escalation"},
               findings={"relevant_event_ids": ["E-0157", "E-0135"], "disputed_event_id": None, "candidate_event_ids": [], "summary": "x"},
               gate={"route": "manual_escalation", "gate_rejected": True, "gate_reasons": ["no draft"]})
    row["checks"]["synth_errors"] = ["synthesize failed: parse"]
    return row


@pytest.fixture
def workbook(ds, tmp_path):
    s01, s07, s02 = ds.scenarios["S-01"], ds.scenarios["S-07"], ds.scenarios["S-02"]
    data = {"category_source": "classify", "split": "dev", "model_id": "claude-sonnet-5-5",
            "golden_path": "dispute_agent/evaluation/golden.xlsx",
            "rows": [normal_row(s01), null_draft_row(s07), {"scenario_id": "S-02", "exception": "KeyError: 'x'"}],
            "totals": {}}
    json_path = tmp_path / "eval_20261027T120000Z.json"
    json_path.write_text(json.dumps(data))
    out = tmp_path / "release.xlsx"
    write_report(load_eval(json_path), json_path, build_cases(ds), out)
    return openpyxl.load_workbook(out)


def table(ws):
    return {r[0]: r[1:] for r in ws.iter_rows(min_row=4, values_only=True) if r[0]}


def test_sheets_and_titles(workbook, ds):
    assert workbook.sheetnames == ["Run", "Cases"] + [f"S-{i:02d}" for i in range(1, 17)]
    run = dict(workbook["Run"].iter_rows(values_only=True))
    assert run["captured_at"] == "2026-10-27 12:00:00Z" and run["model_id"] == "claude-sonnet-5-5"
    assert run["scenarios"] == 3 and run["exceptions"] == 1 and run["langsmith_project"] is None
    ws = workbook["S-01"]
    assert ws["A1"].value == f"S-01 — {ds.scenarios['S-01'].rep_message}"
    assert [c.value for c in ws[3]] == ["field", "golden", "actual"]
    t = table(ws)
    assert t["category"] == ("late_fee_autopay", "late_fee_autopay") and t["route"] == ("explain_events", "explain_events")
    assert t["draft_route"] == (None, "explain_events") and t["disputed_event_id"] == ("E-0117", "E-0117")
    assert t["required_event_ids / cited_event_ids"] == ("E-0081, E-0117", "E-0081, E-0117")
    assert t["existing_case_id / case_id"] == (None, "C-004 (created)") and t["confidence"] == (None, "high")
    assert t["rationale / case_note"] == (ds.scenarios["S-01"].rationale, "The fee stands.")
    assert t["tool_calls_made"] == (None, 3) and t["account_summary"] == (None, "cancelled autopay, fee stands")
    assert "thread_id" not in t
    cases = {r[1]: r for r in workbook["Cases"].iter_rows(min_row=2, values_only=True)}
    assert cases["S-01"][7:] == ("explain_events", "explain_events", "E-0117", "E-0117", "no", "yes", "a1/p1", "OK")
    assert cases["S-04"][2] == "heldout" and len(cases) == 16


def test_null_draft_row(workbook):
    t = table(workbook["S-07"])
    assert t["route"] == ("manual_escalation", "manual_escalation") and t["draft_route"] == (None, NO_DRAFT)
    assert t["rationale / case_note"][1] == NO_DRAFT and t["confidence"] == (None, None) and t["open_questions"] == (None, None)
    assert t["forbidden_event_ids / cited anywhere by the Account Agent"] == ("E-0135", "E-0135")
    assert t["gate_reasons"] == (None, "no draft")
    cases = {r[1]: r for r in workbook["Cases"].iter_rows(min_row=2, values_only=True)}
    assert cases["S-07"][8] == "manual_escalation" and cases["S-07"][11:] == ("yes", "yes", "a1/p1", "OK")


def test_exception_row(workbook, ds):
    t = table(workbook["S-02"])
    assert t["route"] == ("adjustment_review", "EXCEPTION: KeyError: 'x'")
    assert t["rationale / case_note"] == (ds.scenarios["S-02"].rationale, "EXCEPTION: KeyError: 'x'")
    assert t["gate_reasons"] == (None, None)
    cases = {r[1]: r for r in workbook["Cases"].iter_rows(min_row=2, values_only=True)}
    assert cases["S-02"][8] is None and cases["S-02"][-1] == "EXCEPTION"


def test_pending_scenario(workbook, ds):
    t = table(workbook["S-03"])
    assert t["route"] == ("manual_escalation", None) and t["rationale / case_note"] == (ds.scenarios["S-03"].rationale, None)
    assert all(v[1] is None for v in t.values())
    cases = {r[1]: r for r in workbook["Cases"].iter_rows(min_row=2, values_only=True)}
    assert cases["S-03"][-1] == "PENDING" and cases["S-03"][8] is None


def test_rejects_non_report(tmp_path, capsys):
    bad = tmp_path / "account_agent_x.json"
    bad.write_text(json.dumps({"messages": [], "findings": None}))
    assert main(["report", str(bad), "--out", str(tmp_path / "out.xlsx")]) == 2
    assert capsys.readouterr().out.strip() == f"report: {bad} is not an eval report (expected `rows` with scenario_id and `totals`)"
    assert not (tmp_path / "out.xlsx").exists()
    missing = tmp_path / "nope.json"
    assert render(str(missing), str(tmp_path / "out.xlsx")) == 2
