"""The judge: what it is shown, a fake model's scores, the quote rule, retry then error, no draft, and the agreement check."""

import json

import pytest
from pydantic import ValidationError

from dispute_agent.evaluation.golden import read_golden
from dispute_agent.evaluation.judge import (
    JUDGE_AXES, JUDGE_MODEL_ID, MEMO_OPEN, AxisScore, JudgeScores, agreement, agreement_lines, calibrate_check, judge_report,
    judge_row, read_labels, render_judge_input,
)
from dispute_agent.evaluation.report import build_labels
from test_policy_agent import FakeModel
from test_report import normal_row, null_draft_row

CASES = {c.scenario_id: c for c in read_golden()}


def full_scores(score=3, quote="q"):
    return JudgeScores(**{a: AxisScore(score=score, reason="r", quote=quote) for a in JUDGE_AXES})


def row_for(sid, **extra):
    row = normal_row(CASES[sid])
    row["derived"] = {"payment_state": "posted", "refund_state": "not_applicable", "evidence_events": ["E-0081"]}
    row["sync"] = {"issues": [], "rules": [{"section_id": "AP.v1.s4", "condition": "c", "outcome": "explains_outcome", "why": "w", "policy_ambiguous": False}]}
    row["findings"]["other_account_matches"] = []
    return {**row, **extra}


def test_render_blocks_present(ds):
    text = render_judge_input(row_for("S-01"), CASES["S-01"], ds)
    assert text.startswith(f"Complaint: {CASES['S-01'].rep_message}\nCategory: late_fee_autopay\nExpected route (answer key): explain_events")
    assert f"Golden rationale (answer key): {CASES['S-01'].rationale}" in text and "Account memo: none" in text
    assert '"event_id": "E-0081"' in text and "- AP.v1.s4 | explains_outcome | ambiguous=false | condition: c | why: w" in text
    assert '"case_note": "The fee stands."' in text and "Final route (the draft's, or manual_escalation with no draft): explain_events" in text
    assert "Derived facts:" in text and "evidence_events" not in text.split("Derived facts:")[1].split("\n")[0]
    s04 = render_judge_input(row_for("S-04"), CASES["S-04"], ds)
    assert f"Account memo: {MEMO_OPEN} " in s04  # the injection memo is shown as data, inside the delimiters
    unknown = render_judge_input(row_for("S-01", derived={"evidence_events": ["E-9999"]}), CASES["S-01"], ds)
    assert '"note": "not in the dataset"' in unknown


def test_fake_model_scores(ds):
    model = FakeModel(full_scores(2))
    out = judge_row(row_for("S-01"), CASES["S-01"], ds, model, config={"tags": ["t"]})
    assert out == {a: {"score": 2, "reason": "r", "quote": "q"} for a in JUDGE_AXES} and model.calls == 1
    as_dict = FakeModel(full_scores(3, "").model_dump())
    assert judge_row(row_for("S-01"), CASES["S-01"], ds, as_dict)["no_leaks_or_promises"] == {"score": 3, "reason": "r", "quote": ""}


def test_quote_required_below_3():
    with pytest.raises(ValidationError, match="quote is required"):
        AxisScore(score=2, reason="x")
    assert AxisScore(score=3, reason="x").quote == "" and AxisScore(score=0, reason="x", quote="bad span").score == 0
    with pytest.raises(ValidationError):
        AxisScore(score=4, reason="x", quote="q")


def test_retry_then_error(ds):
    model = FakeModel(RuntimeError("boom"))
    out = judge_row(row_for("S-01"), CASES["S-01"], ds, model)
    assert out == {"judge_error": "RuntimeError: boom | RuntimeError: boom"} and model.calls == 2
    bad = FakeModel({a: {"score": 1, "reason": "r", "quote": ""} for a in JUDGE_AXES})  # below 3 with no quote: schema rejection
    assert "judge_error" in judge_row(row_for("S-01"), CASES["S-01"], ds, bad) and bad.calls == 2


def test_null_draft_not_called(ds):
    model = FakeModel(full_scores())
    out = judge_row(null_draft_row(CASES["S-07"]), CASES["S-07"], ds, model)
    assert out == {a: {"score": 0, "reason": "no draft", "quote": ""} for a in JUDGE_AXES} and model.calls == 0
    out = judge_row({"scenario_id": "S-02", "exception": "KeyError"}, CASES["S-02"], ds, model)
    assert out["no_invented_facts"]["reason"] == "no draft (exception)" and model.calls == 0


def test_judge_report_writes_sidecar(ds, tmp_path, capsys):
    data = {"rows": [row_for("S-01"), null_draft_row(CASES["S-07"]), {"scenario_id": "S-02", "exception": "KeyError"}]}
    out = judge_report(tmp_path / "eval_x.json", data, ds, list(CASES.values()), FakeModel(full_scores(2)), out_dir=tmp_path)
    side = json.loads(out.read_text())
    assert out.name.startswith("judge_") and side["eval_path"] == str(tmp_path / "eval_x.json") and side["judge_model"] == JUDGE_MODEL_ID
    assert side["rows"]["S-01"]["no_leaks_or_promises"]["score"] == 2 and side["rows"]["S-07"]["no_leaks_or_promises"]["score"] == 0 and side["rows"]["S-02"]["no_leaks_or_promises"]["score"] == 0
    assert "judged 3 scenarios, 0 errors" in capsys.readouterr().out


def judge_rows(**per_sid):
    return {sid: {a: {"score": s, "reason": "r", "quote": "q"} for a in JUDGE_AXES} for sid, s in per_sid.items()}


def test_agreement_math():
    judge = judge_rows(**{"S-01": 3, "S-02": 1, "S-03": 2}) | {"S-05": {"judge_error": "boom"}}
    labels = {sid: {a: h for a in JUDGE_AXES} for sid, h in (("S-01", 3), ("S-02", 3), ("S-03", 1), ("S-05", 2))}
    ag = agreement(judge, labels)
    a = ag["no_invented_facts"]
    assert a["n"] == 3 and a["exact"] == 0.33 and a["within_1"] == 0.67 and a["judge_mean"] == 2.0 and a["human_mean"] == 2.33
    assert a["disagreements"] == [("S-02", 1, 3), ("S-03", 2, 1)]
    assert agreement_lines(ag)[1].split() == ["no_invented_facts", "3", "0.33", "0.67", "2.0", "2.33"]


def test_partial_labels():
    judge = judge_rows(**{"S-01": 3, "S-02": 2})
    labels = {"S-01": {a: 3 for a in JUDGE_AXES}, "S-02": {"no_invented_facts": None, "matches_rationale": 2, "ready_to_file": None, "no_leaks_or_promises": 2}}
    ag = agreement(judge, labels)
    assert ag["no_invented_facts"]["n"] == 1 and ag["matches_rationale"]["n"] == 2 and ag["ready_to_file"]["within_1"] == 1.0
    assert agreement(judge, {})["no_leaks_or_promises"] == {"n": 0, "exact": None, "within_1": None, "judge_mean": None, "human_mean": None, "disagreements": []}


def test_labels_build_check_and_exit_codes(ds, tmp_path, capsys):
    data = {"rows": [row_for("S-01"), null_draft_row(CASES["S-07"]), row_for("S-04")]}  # S-04 is held out: never labelled
    labels_path = tmp_path / "labels.xlsx"
    assert build_labels(data, list(CASES.values()), labels_path) == 2
    labels = read_labels(labels_path)
    assert list(labels) == ["S-01", "S-07"] and labels["S-01"] == {a: None for a in JUDGE_AXES}
    judge_path = tmp_path / "judge.json"
    judge_path.write_text(json.dumps({"eval_path": "x", "judge_model": "m", "rows": judge_rows(**{"S-01": 3, "S-07": 0})}))
    assert calibrate_check(judge_path, labels_path) == 2 and "no labels filled" in capsys.readouterr().out
    import openpyxl
    wb = openpyxl.load_workbook(labels_path)
    ws = wb["Labels"]
    header = [c.value for c in ws[1]]
    for r, scores in ((2, (3, 3, 3, 3)), (3, (3, 0, 0, 1))):
        for axis, s in zip(JUDGE_AXES, scores):
            ws.cell(row=r, column=header.index(axis) + 1).value = s
    wb.save(labels_path)
    assert calibrate_check(judge_path, labels_path) == 0  # n < 6 per axis: no axis can fail the threshold
    out = capsys.readouterr().out
    assert "below 75% within-1: none" in out and "matches_rationale" in out
