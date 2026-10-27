"""The golden workbook: a field-for-field copy of scenarios.xlsx, round-trips, blanks stay blank, drift is caught."""

import openpyxl

from dispute_agent.evaluation.evaluate import select_scenarios
from dispute_agent.evaluation.evaluate import score
from dispute_agent.evaluation.golden import (
    GOLDEN_XLSX, PRODUCERS, Case, build_cases, check, empty_stub, golden_echo, read_golden, write_golden,
)
from dispute_agent.evaluation.metrics import AXES, axis_verdicts, grade

SHEETS = ["Cases"] + [f"S-{i:02d}" for i in range(1, 17)]


def test_sixteen_cases_match_scenarios(ds):
    cases = build_cases(ds)
    assert len(cases) == 16 and all(isinstance(c, Case) for c in cases)
    for c in cases:
        assert c.model_dump() == ds.scenarios[c.scenario_id].model_dump()
    assert set(Case.model_fields) == set(type(ds.scenarios["S-01"]).model_fields)  # no golden field beyond scenarios.xlsx


def test_round_trip(ds, tmp_path):
    path = tmp_path / "golden.xlsx"
    write_golden(build_cases(ds), path)
    assert openpyxl.load_workbook(path, read_only=True).sheetnames == SHEETS
    back = read_golden(path)
    assert [c.model_dump() for c in back] == [s.model_dump() for s in ds.scenarios.values()]


def test_blanks_are_empty_cells(ds, tmp_path):
    path = tmp_path / "golden.xlsx"
    write_golden(build_cases(ds), path)
    wb = openpyxl.load_workbook(path, read_only=True)
    s06 = {r[0]: r[1] for r in wb["S-06"].iter_rows(values_only=True) if r[0]}
    assert s06["expected_disputed_event_id"] is None and s06["expected_existing_case_id"] is None
    assert s06["expected_required_event_ids"] == "E-0109" and s06["expected_forbidden_event_ids"] == "E-0158"
    s01 = {r[0]: r[1] for r in wb["S-01"].iter_rows(values_only=True) if r[0]}
    assert s01["expected_forbidden_event_ids"] is None and s01["expected_required_event_ids"] == "E-0081, E-0117"
    s11 = {r[0]: r[1] for r in wb["S-11"].iter_rows(values_only=True) if r[0]}
    assert s11["expected_existing_case_id"] == "C-002"
    for ws in wb.worksheets:
        assert not any(v == "None" or v == "[]" for row in ws.iter_rows(values_only=True) for v in row)
    index = list(wb["Cases"].iter_rows(values_only=True))
    assert index[0][:3] == ("#", "scenario_id", "split") and len(index) == 17
    assert index[6][1] == "S-06" and index[6][9] is None  # expected_disputed_event_id blank, not "None"


def test_check_detects_drift(ds, tmp_path):
    assert check(ds, GOLDEN_XLSX) == []  # the committed workbook is in sync (acceptance criterion 2)
    path = tmp_path / "golden.xlsx"
    write_golden(build_cases(ds), path)
    assert check(ds, path) == []
    wb = openpyxl.load_workbook(path)
    ws = wb["S-01"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "expected_route":
            row[1].value = "manual_escalation"
    wb.save(path)
    diffs = check(ds, path)
    assert len(diffs) == 1 and diffs[0].startswith("S-01.expected_route: workbook 'manual_escalation'")
    assert check(ds, tmp_path / "missing.xlsx") == [f"{tmp_path / 'missing.xlsx'} is missing; run --build"]


def test_check_detects_index_drift(ds, tmp_path):
    path = tmp_path / "golden.xlsx"
    write_golden(build_cases(ds), path)
    wb = openpyxl.load_workbook(path)
    wb["Cases"]["I2"].value = "adjustment_review"  # S-01's expected_route in the index only
    wb.save(path)
    assert check(ds, path) == ["Cases S-01.expected_route: workbook 'adjustment_review', scenarios.xlsx 'explain_events'"]


def test_split_all_selects_sixteen(ds):
    assert len(select_scenarios(ds, "all", None)) == 16
    assert [s.scenario_id for s in select_scenarios(ds, "dev", None)] == [
        f"S-{i:02d}" for i in (1, 2, 3, 5, 6, 7, 9, 10, 11, 13, 14, 15)
    ]
    assert [s.scenario_id for s in select_scenarios(ds, "heldout", None)] == ["S-04", "S-08", "S-12", "S-16"]
    assert [s.scenario_id for s in select_scenarios(ds, "all", "S-04")] == ["S-04"]


def test_golden_echo_all_pass_16(ds, sections):
    """The scorer's calibration: the key itself passes every axis on every scenario, S-11 and S-15 included."""
    for case in read_golden():
        inv = golden_echo(ds, sections, case)
        verdicts = axis_verdicts(score(case, inv, ds))
        assert verdicts == {a: "PASS" for a in AXES}, (case.scenario_id, verdicts)
        assert inv.case.created is (case.expected_existing_case_id is None)


def test_empty_stub_fails(ds, sections):
    """An empty answer never grades PASS; its fallback route is right only where the key wants manual_escalation."""
    for case in read_golden():
        verdicts = axis_verdicts(score(case, empty_stub(ds, sections, case), ds))
        assert grade(verdicts) == "FAIL" and verdicts["evidence"] == "FAIL" and verdicts["confidence"] == "FAIL"
        assert verdicts["route"] == ("PARTIAL" if case.expected_route == "manual_escalation" else "FAIL")
    assert set(PRODUCERS) == {"golden_echo", "empty_stub"}
