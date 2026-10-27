"""One eval JSON plus golden.xlsx -> a release workbook: `Run`, `Cases`, `Verdict Matrix`, `Slices`, one `S-xx` sheet per
golden case; with --judge a judge panel per sheet, with --before a `Compare` sheet, with --labels a `Calibration` sheet.
Every actual value is copied from the JSON row as the run produced it; verdicts come from metrics.py, scores from the
judge sidecar. Nothing is scored here. `build_labels` writes the workbook Rachit scores by hand (`calibrate --build`).
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from dispute_agent.evaluation.evaluate import confusion_lines
from dispute_agent.evaluation.golden import GOLDEN_XLSX, HEADER_FILL, HEADER_FONT, WRAP, Case, cell, read_golden
from dispute_agent.evaluation.judge import SCALE_TEXT, agreement, read_labels
from dispute_agent.evaluation.metrics import (
    AXES, JUDGE_AXES, SLICE_KEYS, compare, directional, grade, judge_mean, judge_scores, judge_verdict, pass_counts,
    row_verdicts, score_str, slice_table,
)

RUN_KEYS = ("json_path", "captured_at", "model_id", "producer", "split", "scenarios", "exceptions", "golden_path", "judge_path", "before_path")
CASE_COLS = [
    "#", "scenario_id", "split", "trap", "wording", "note_driven", "complaint", "golden_route", "actual_route",
    "golden_disputed", "actual_disputed", "policy_consistent", "case_created", "status",
]
MATRIX_COLS = ["Q", "Trap", "Wording", "Note", *AXES, "Grade", "Score", *JUDGE_AXES, "Judge mean", "P1-P7"]
LABEL_COLS = ["scenario_id", "trap", "complaint", "expected_route", "rationale", "actual_route", "case_note", "open_questions", *JUDGE_AXES, "notes"]
FORBIDDEN_CITED = "forbidden_event_ids / cited anywhere by the Account Agent"
FIELD_AXIS = {  # which axis governs each golden field's verdict cell
    "category": "guardrails", "route": "route", "disputed_event_id": "disputed_event", "confidence": "confidence",
    "required_event_ids / cited_event_ids": "evidence", FORBIDDEN_CITED: "evidence", "existing_case_id / case_id": "guardrails",
    "policy_ids / policy_source_ids": "policy",
}
NO_DRAFT = "(no draft)"
FILLS = {"PASS": "C6EFCE", "PARTIAL": "FFEB9C", "FAIL": "FFC7CE", "ERR": "D9D9D9", "improved": "C6EFCE", "regressed": "FFC7CE"}


def load_eval(path: Path) -> dict[str, Any]:
    """The JSON an `eval` run wrote; anything else (an `ask` dump, a checkpoint) raises ValueError."""
    data = json.loads(path.read_text())
    rows = data.get("rows") if isinstance(data, dict) else None
    if not isinstance(rows, list) or "totals" not in data or not all("scenario_id" in r for r in rows):
        raise ValueError(f"{path} is not an eval report (expected `rows` with scenario_id and `totals`)")
    return data


def load_judge(path: Path, json_path: Path) -> dict[str, Any]:
    """The `judge` sidecar for exactly this eval JSON; a sidecar of another run raises ValueError."""
    data = json.loads(path.read_text())
    if not isinstance(data, dict) or "rows" not in data or "eval_path" not in data:
        raise ValueError(f"{path} is not a judge sidecar")
    if Path(data["eval_path"]).resolve() != json_path.resolve():
        raise ValueError(f"{path} judged {data['eval_path']}, not {json_path}")
    return data


def captured_at(path: Path) -> str:
    m = re.search(r"(\d{4})(\d{2})(\d{2})T(\d{2})(\d{2})(\d{2})Z", path.name)
    return f"{m[1]}-{m[2]}-{m[3]} {m[4]}:{m[5]}:{m[6]}Z" if m else ""


def yes(v: Any) -> str:
    return "yes" if v else "no"


def consistent(row: dict[str, Any]) -> str | None:
    """`policy_consistent` as ok / ruled out; None for reports made before the check existed (v0.8)."""
    v = (row.get("checks") or {}).get("policy_consistent")
    return None if v is None else "ok" if v else "ruled out"


def fill(c, verdict: Any) -> None:
    if verdict in FILLS:
        c.fill = PatternFill("solid", fgColor=FILLS[verdict])


def cited_by_account_agent(row: dict[str, Any]) -> set[str]:
    """Every event ID the Account Agent's findings mention: relevant, candidate and disputed, the set score() checks."""
    f = row.get("findings") or {}
    return set(f.get("relevant_event_ids") or []) | set(f.get("candidate_event_ids") or []) | ({f.get("disputed_event_id")} - {None})


def actual_fields(case: Case, row: dict[str, Any]) -> dict[str, Any]:
    """The `actual` column of a scenario sheet, keyed by the golden field it sits beside."""
    draft, checks, derived, filed = row.get("draft"), row.get("checks") or {}, row.get("derived") or {}, row.get("case") or {}
    forbidden_cited = sorted(set(case.expected_forbidden_event_ids) & cited_by_account_agent(row))
    return {
        "category": row.get("category_pred"),
        "route": checks.get("route_final"),
        "draft_route": draft["review_route"] if draft else NO_DRAFT,
        "disputed_event_id": (row.get("findings") or {}).get("disputed_event_id"),
        "required_event_ids / cited_event_ids": cell(draft["event_ids"]) if draft else NO_DRAFT,
        FORBIDDEN_CITED: cell(forbidden_cited),
        "payment_state": derived.get("payment_state"),
        "refund_state": derived.get("refund_state"),
        "existing_case_id / case_id": f"{filed['case_id']} ({'created' if filed.get('created') else 'existing'})" if filed else None,
        "policy_ids / policy_source_ids": cell(draft["policy_source_ids"]) if draft else NO_DRAFT,
        "confidence": draft["confidence"] if draft else None,
        "open_questions": ("\n".join(draft.get("open_questions") or []) or None) if draft else None,
        "rationale / case_note": draft["case_note"] if draft else NO_DRAFT,
    }


def golden_fields(case: Case) -> dict[str, Any]:
    return {
        "category": case.expected_category, "route": case.expected_route, "draft_route": None,
        "disputed_event_id": case.expected_disputed_event_id,
        "required_event_ids / cited_event_ids": cell(case.expected_required_event_ids),
        FORBIDDEN_CITED: cell(case.expected_forbidden_event_ids),
        "payment_state": case.expected_payment_state, "refund_state": case.expected_refund_state,
        "existing_case_id / case_id": case.expected_existing_case_id,
        "policy_ids / policy_source_ids": cell(case.expected_policy_ids),
        "confidence": None, "open_questions": None, "rationale / case_note": case.rationale,
    }


def run_detail(row: dict[str, Any]) -> list[tuple[str, Any]]:
    checks, sync = row.get("checks") or {}, row.get("sync") or {}
    return [
        ("synth_errors", cell(checks.get("synth_errors"))), ("sync_issues", cell(sync.get("issues"))),
        ("account_summary", (row.get("findings") or {}).get("summary")), ("tool_calls_made", checks.get("tool_calls")),
        ("tool_errors", cell(row.get("tool_errors"))),
    ]


def _append_header(ws, values: list[Any]) -> None:
    ws.append(values)
    for c in ws[ws.max_row]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL


def write_scenario_sheet(wb: Workbook, case: Case, row: dict[str, Any] | None, judge: dict[str, Any] | None) -> None:
    ws = wb.create_sheet(title=case.scenario_id)
    ws.append([f"{case.scenario_id} — {case.rep_message}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    _append_header(ws, ["field", "golden", "actual", "verdict"])
    golden = golden_fields(case)
    if row is None:
        actual: dict[str, Any] = {}
    elif "exception" in row:
        actual = {k: f"EXCEPTION: {row['exception']}" for k in golden}
    else:
        actual = actual_fields(case, row)
    verdicts = row_verdicts(row) if row is not None else {}
    for name, g in golden.items():
        if name == "rationale / case_note":
            ws.append([])
        verdict = verdicts.get(FIELD_AXIS.get(name, ""))
        ws.append([name, g, actual.get(name), verdict])
        fill(ws.cell(row=ws.max_row, column=4), verdict)
    ws.append([])
    _append_header(ws, ["Run detail"])
    for name, value in run_detail(row or {}):
        ws.append([name, None, value if row and "exception" not in row else None])
    if judge is not None:
        ws.append([])
        _append_header(ws, ["Judge", "score", "reason", "quote"])
        judged = judge["rows"].get(case.scenario_id)
        if judged is None:
            ws.append(["not judged"])
        elif "judge_error" in judged:
            ws.append(["ERR", None, judged["judge_error"]])
            fill(ws.cell(row=ws.max_row, column=1), "ERR")
        else:
            for axis in JUDGE_AXES:
                ws.append([axis, judged[axis]["score"], judged[axis]["reason"], judged[axis]["quote"]])
                fill(ws.cell(row=ws.max_row, column=2), judge_verdict(judged[axis]["score"]))
    for col, width in zip("ABCD", (28, 60, 60, 50)):
        ws.column_dimensions[col].width = width
    for r in ws.iter_rows(min_row=3):
        for c in r:
            c.alignment = WRAP
    ws.freeze_panes = "A4"


def case_row(n: int, case: Case, row: dict[str, Any] | None) -> list[Any]:
    base = [n, case.scenario_id, case.split, case.trap, case.wording, case.note_driven, case.rep_message[:80], case.expected_route]
    if row is None:
        return base + [None, case.expected_disputed_event_id, None, None, None, "PENDING"]
    if "exception" in row:
        return base + [None, case.expected_disputed_event_id, None, None, None, "EXCEPTION"]
    filed = row.get("case") or {}
    return base + [
        (row.get("checks") or {}).get("route_final"), case.expected_disputed_event_id,
        (row.get("findings") or {}).get("disputed_event_id"), consistent(row), yes(filed.get("created")), "OK",
    ]


def matrix_row(case: Case, row: dict[str, Any] | None, judge: dict[str, Any] | None) -> list[Any]:
    base = [case.scenario_id, case.trap, case.wording, yes(case.note_driven)]
    if row is None:
        return base + [None] * len(AXES) + ["PENDING", None] + [None] * (len(JUDGE_AXES) + 2)
    verdicts = row_verdicts(row)
    scores = judge_scores(judge["rows"].get(case.scenario_id)) if judge else {a: None for a in JUDGE_AXES}
    judge_cells = [("ERR" if judge and s is None else s) for s in scores.values()]
    cons = "—" if "exception" in row else consistent(row)
    return base + [verdicts[a] for a in AXES] + [grade(verdicts), score_str(verdicts)] + judge_cells + [
        judge_mean(scores) if judge else None, cons,
    ]


def write_matrix(wb: Workbook, cases: list[Case], rows: dict[str, dict[str, Any]], judge: dict[str, Any] | None) -> None:
    ws = wb.create_sheet(title="Verdict Matrix")
    _append_header(ws, MATRIX_COLS)
    for case in cases:
        ws.append(matrix_row(case, rows.get(case.scenario_id), judge))
        for c in ws[ws.max_row]:  # verdict strings fill as themselves; a 0-3 judge score fills by its verdict
            fill(c, judge_verdict(c.value) if isinstance(c.value, int) and not isinstance(c.value, bool) else c.value)
    present = [rows[c.scenario_id] for c in cases if c.scenario_id in rows]
    counts = pass_counts(present)
    judged = [judge_scores(judge["rows"].get(r["scenario_id"])) for r in present] if judge else []
    mean = lambda v: round(sum(v) / len(v), 2) if v else None  # noqa: E731
    axis_means = [mean([j[axis] for j in judged if j[axis] is not None]) for axis in JUDGE_AXES]
    _append_header(ws, ["total", None, None, None] + ["%d/%d" % counts[a] for a in AXES] + ["%d/%d" % counts["grade"], None]
                   + axis_means + [mean([m for m in map(judge_mean, judged) if m is not None]), None])
    ws.append([])
    _append_header(ws, ["expected \\ final route"])
    for line in confusion_lines([r for r in present if "checks" in r])[1:]:
        ws.append(line.split())
    ws.freeze_panes = "B2"


def write_slices(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title="Slices")
    table = slice_table(rows)
    for key in SLICE_KEYS:
        _append_header(ws, [key, *AXES, "grade"])
        for value, cells in table[key].items():
            n = cells["grade"][1]
            ws.append([f"{value} (directional)" if directional(n) else value] + [f"{p}/{t}" for p, t in cells.values()])
        ws.append([])
    ws.column_dimensions["A"].width = 32


def write_compare(wb: Workbook, before: dict[str, Any], before_path: Path, after: dict[str, Any], after_path: Path) -> None:
    ws = wb.create_sheet(title="Compare")
    c = compare(before["rows"], after["rows"])
    for k, v in (("before", str(before_path)), ("after", str(after_path)), ("scenarios compared", len(c["shared"])),
                 ("excluded (in one run only)", ", ".join(c["excluded"]) or "none")):
        ws.append([k, v])
    ws.append([])
    _append_header(ws, ["axis", "before PASS", "after PASS"])
    for name in AXES + ("grade",):
        ws.append([name, "%d/%d" % c["before"][name], "%d/%d" % c["after"][name]])
    ws.append([])
    _append_header(ws, ["scenario", "axis", "before", "after", "Δ"])
    for sid, axis, vb, va, delta in c["diffs"]:
        ws.append([sid, axis, vb, va, delta])
        for col, v in ((3, vb), (4, va), (5, delta)):
            fill(ws.cell(row=ws.max_row, column=col), v)
    if not c["diffs"]:
        ws.append(["no changes"])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 70


def write_calibration(wb: Workbook, judge: dict[str, Any], labels: dict[str, dict[str, int | None]]) -> None:
    ws = wb.create_sheet(title="Calibration")
    ag = agreement(judge["rows"], labels)
    if not any(a["n"] for a in ag.values()):
        ws.append(["no labels yet"])
        return
    _append_header(ws, ["axis", "n", "exact", "within-1", "judge mean", "human mean"])
    for axis, a in ag.items():
        ws.append([axis, a["n"], a["exact"], a["within_1"], a["judge_mean"], a["human_mean"]])
    ws.append([])
    _append_header(ws, ["scenario", "axis", "judge", "human"])
    for axis, a in ag.items():
        for sid, j, h in a["disagreements"]:
            ws.append([sid, axis, j, h])
    ws.column_dimensions["A"].width = 24


def build_labels(data: dict[str, Any], cases: list[Case], out: Path) -> int:
    """The workbook Rachit scores by hand: the 12 dev notes, four empty 0-3 columns. Returns the row count."""
    rows = {r["scenario_id"]: r for r in data["rows"]}
    wb = Workbook()
    ws = wb.active
    ws.title = "Labels"
    _append_header(ws, LABEL_COLS)
    n = 0
    for case in cases:
        row = rows.get(case.scenario_id)
        if case.split != "dev" or row is None:
            continue
        draft, route = row.get("draft"), (row.get("checks") or {}).get("route_final")
        ws.append([case.scenario_id, case.trap, case.rep_message, case.expected_route, case.rationale, route,
                   draft["case_note"] if draft else NO_DRAFT, "\n".join(draft.get("open_questions") or []) if draft else None])
        n += 1
    dv = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    ws.add_data_validation(dv)
    first, last = LABEL_COLS.index(JUDGE_AXES[0]) + 1, LABEL_COLS.index(JUDGE_AXES[-1]) + 1
    dv.add(f"{ws.cell(row=2, column=first).column_letter}2:{ws.cell(row=n + 1, column=last).column_letter}{n + 1}")
    for col, width in zip("ABCDEFGH", (12, 14, 50, 20, 60, 20, 60, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B2"
    how = wb.create_sheet(title="How to score")
    for line in SCALE_TEXT.splitlines():
        how.append([line])
    how.column_dimensions["A"].width = 120
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return n


def write_report(
    data: dict[str, Any], json_path: Path, cases: list[Case], out: Path, judge: dict[str, Any] | None = None,
    before: tuple[dict[str, Any], Path] | None = None, labels: dict[str, dict[str, int | None]] | None = None,
) -> None:
    rows = {r["scenario_id"]: r for r in data["rows"]}
    wb = Workbook()
    run = wb.active
    run.title = "Run"
    values = {
        "json_path": str(json_path), "captured_at": captured_at(json_path), "model_id": data.get("model_id", ""),
        "producer": data.get("producer", "agent"), "split": data.get("split", ""), "scenarios": len(rows),
        "exceptions": sum("exception" in r for r in rows.values()), "golden_path": data.get("golden_path", ""),
        "judge_path": f"{judge['judge_model']} via judge sidecar" if judge else "", "before_path": str(before[1]) if before else "",
    }
    for key in RUN_KEYS:
        run.append([key, values[key]])
    run.column_dimensions["A"].width, run.column_dimensions["B"].width = 20, 70
    idx = wb.create_sheet(title="Cases")
    _append_header(idx, CASE_COLS)
    for n, case in enumerate(cases, 1):
        idx.append(case_row(n, case, rows.get(case.scenario_id)))
    idx.freeze_panes = "A2"
    write_matrix(wb, cases, rows, judge)
    meta = {c.scenario_id: {"trap": c.trap, "wording": c.wording, "note_driven": c.note_driven, "category_expected": c.expected_category, "route_expected": c.expected_route} for c in cases}
    write_slices(wb, [meta[sid] | rows[sid] for sid in meta if sid in rows])  # an exception row carries no tags; the golden case does
    if before:
        write_compare(wb, before[0], before[1], data, json_path)
    if labels is not None and judge is not None:
        write_calibration(wb, judge, labels)
    for case in cases:
        write_scenario_sheet(wb, case, rows.get(case.scenario_id), judge)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def render(json_path: str, out: str, golden_path: Path = GOLDEN_XLSX, judge: str | None = None,
           before: str | None = None, labels: str | None = None) -> int:
    """The `report` command: exit 0 on success, 2 when an input is not what its flag says."""
    try:
        if labels and not judge:
            raise ValueError("--labels needs --judge (calibration compares the judge to the labels)")
        data = load_eval(Path(json_path))
        judge_data = load_judge(Path(judge), Path(json_path)) if judge else None
        before_data = (load_eval(Path(before)), Path(before)) if before else None
        label_data = read_labels(Path(labels)) if labels else None
    except (ValueError, json.JSONDecodeError, OSError, KeyError) as exc:
        print(f"report: {exc}")
        return 2
    cases = read_golden(golden_path)
    write_report(data, Path(json_path), cases, Path(out), judge_data, before_data, label_data)
    extras = [name for name, flag in (("judge", judge), ("compare", before), ("calibration", labels)) if flag]
    print(f"wrote {out}: {len(cases)} scenario sheets, {len(data['rows'])} with a run" + (f", with {', '.join(extras)}" if extras else ""))
    return 0
