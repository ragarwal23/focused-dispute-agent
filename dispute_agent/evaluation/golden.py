"""The golden workbook: data/scenarios.xlsx copied into a `Cases` index plus one sheet per scenario.

golden.xlsx is generated, committed, and drift-checked; the release report reads it, never scenarios.xlsx.
A `Case` has exactly the fields of a `Scenario`, so `score()` and the report accept either.

    python -m dispute_agent.evaluation.golden --build    rewrite golden.xlsx from data/scenarios.xlsx
    python -m dispute_agent.evaluation.golden --check    exit 1 if golden.xlsx differs from a rebuild

The producers at the bottom turn a golden case into an `Investigation` with no LLM: `golden_echo` is the key itself
(the scorer must pass it on every scenario), `empty_stub` is an agent that found nothing (the scorer must fail it).
`eval --producer` runs them through the same scoring and JSON as a live run.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from dispute_agent.account_agent.agent import AccountFindings, OtherAccountMatch
from dispute_agent.account_agent.derive import derive_states
from dispute_agent.data import Dataset, Scenario, load_dataset
from dispute_agent.orchestrator.run import Investigation
from dispute_agent.policy_agent.agent import PolicyFindings, RuleReading
from dispute_agent.policy_agent.retrieve import PolicySection, RetrievedSection
from dispute_agent.review.create_review_case import create_review_case
from dispute_agent.sync.sync import sync
from dispute_agent.sync.synthesize import Draft, SynthesisResult

GOLDEN_XLSX = Path(__file__).resolve().parent / "golden.xlsx"
GOLDEN_PATH = "dispute_agent/evaluation/golden.xlsx"  # as written into every eval JSON
INDEX_SHEET = "Cases"
INDEX_COLS = [
    "#", "scenario_id", "split", "account_id", "wording", "note_driven", "trap", "expected_category",
    "expected_route", "expected_disputed_event_id", "rep_message", "rationale",
]
IDENTITY_FIELDS = ("scenario_id", "split", "account_id", "wording", "note_driven", "trap")  # read from the index
SHEET_FIELDS = (  # the per-scenario sheet, in this order
    "expected_category", "expected_route", "expected_disputed_event_id", "expected_required_event_ids",
    "expected_forbidden_event_ids", "expected_payment_state", "expected_refund_state", "expected_existing_case_id",
    "expected_policy_ids",
)
TEXT_FIELDS = ("rep_message", "rationale")
LIST_FIELDS = {"expected_required_event_ids", "expected_forbidden_event_ids", "expected_policy_ids"}

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(wrap_text=True, vertical="top")


class Case(Scenario):
    """One golden case: a Scenario read back from golden.xlsx (same fields, same list parsing)."""


def build_cases(ds: Dataset) -> list[Case]:
    return [Case(**s.model_dump()) for s in ds.scenarios.values()]


def cell(value: Any) -> Any:
    """A model value as a cell: lists joined like scenarios.xlsx, blanks as empty cells, never 'None'."""
    if value is None or value == []:
        return None
    if isinstance(value, list):
        return ", ".join(value)
    return value


def style_header(ws, row: int = 1) -> None:
    for c in ws[row]:
        c.font, c.fill = HEADER_FONT, HEADER_FILL


def write_golden(cases: list[Case], path: Path = GOLDEN_XLSX) -> None:
    wb = Workbook()
    idx = wb.active
    idx.title = INDEX_SHEET
    idx.append(INDEX_COLS)
    style_header(idx)
    for n, case in enumerate(cases, 1):
        idx.append([n] + [cell(getattr(case, col)) for col in INDEX_COLS[1:]])
        ws = wb.create_sheet(title=case.scenario_id)
        ws.append(["field", "golden"])
        style_header(ws)
        for name in SHEET_FIELDS:
            ws.append([name, cell(getattr(case, name))])
        ws.append([])
        for name in TEXT_FIELDS:
            ws.append([name, getattr(case, name)])
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 60
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
        ws.freeze_panes = "A2"
    idx.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _sheet_values(ws) -> dict[str, Any]:
    """`field | golden` rows as a dict; blank fields and blank rows are skipped."""
    out = {}
    for name, value, *_ in ws.iter_rows(min_row=2, values_only=True):
        if name:
            out[str(name)] = value
    return out


def read_golden(path: Path = GOLDEN_XLSX) -> list[Case]:
    """Parse the committed workbook. Stops at the first row with no scenario_id (max_row is stale on these files)."""
    wb = load_workbook(path, data_only=True)
    try:
        rows = wb[INDEX_SHEET].iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        cases = []
        for raw in rows:
            row = dict(zip(header, raw))
            if not row.get("scenario_id"):
                break
            values = {name: row[name] for name in IDENTITY_FIELDS}
            values.update(_sheet_values(wb[row["scenario_id"]]))
            cases.append(Case(**values))
        return cases
    finally:
        wb.close()


def check(ds: Dataset, path: Path = GOLDEN_XLSX) -> list[str]:
    """Field-level differences between the committed workbook and a rebuild; empty means in sync."""
    if not path.exists():
        return [f"{path} is missing; run --build"]
    built = {c.scenario_id: c.model_dump() for c in build_cases(ds)}
    stored = {c.scenario_id: c.model_dump() for c in read_golden(path)}
    diffs = [f"{sid}: only in {'workbook' if sid in stored else 'scenarios.xlsx'}" for sid in built.keys() ^ stored.keys()]
    for sid in built.keys() & stored.keys():
        for name, want in built[sid].items():
            if stored[sid][name] != want:
                diffs.append(f"{sid}.{name}: workbook {stored[sid][name]!r}, scenarios.xlsx {want!r}")
    return diffs + _index_diffs(ds, path)


def _index_diffs(ds: Dataset, path: Path) -> list[str]:
    """The Cases index repeats some sheet fields for reading; a hand edit there must fail --check too."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[INDEX_SHEET].iter_rows(values_only=True)
        header = [str(h) for h in next(rows)]
        diffs = []
        for raw in rows:
            row = dict(zip(header, raw))
            s = ds.scenarios.get(row.get("scenario_id"))
            if s is None:
                break
            for name in INDEX_COLS[2:]:
                if row.get(name) != cell(getattr(s, name)):
                    diffs.append(f"Cases {s.scenario_id}.{name}: workbook {row.get(name)!r}, scenarios.xlsx {cell(getattr(s, name))!r}")
        return diffs
    finally:
        wb.close()


def key_findings(case: Scenario) -> AccountFindings:
    """What a perfect Account Agent would return for this scenario, built from the answer key."""
    other = OtherAccountMatch(event_id="E-0135", account_id="A-1008", reference="PHN-0048213", why="keyed to another customer")
    return AccountFindings(
        relevant_event_ids=case.expected_required_event_ids,
        disputed_event_id=case.expected_disputed_event_id,
        candidate_event_ids=case.expected_required_event_ids if case.trap == "ambiguous" else [],
        other_account_matches=[other] if case.scenario_id == "S-07" else [],
        claimed_but_not_found=["per key"] if case.trap in ("missing_record", "wrong_account") else [],
        notes_summary="", summary="",
    )


def investigation(
    ds: Dataset, sections: dict[str, PolicySection], case: Scenario, findings: AccountFindings | None,
    draft: Draft | None,
) -> Investigation:
    """An Investigation built from the answer key around the given findings and draft; code steps run for real."""
    account = {"findings": findings, "budget_hit": False, "tool_errors": [], "tool_calls_made": 2, "messages": []}
    derived = derive_states(findings, case.expected_category, case.account_id, ds)
    policy = PolicyFindings(
        policy_date=derived.policy_date, category=case.expected_category,
        retrieved=[RetrievedSection(section=sections[i], rank=n, score=1.0) for n, i in enumerate(case.expected_policy_ids, 1) if i in sections],
        rules=[RuleReading(section_id=i, condition="per key", outcome="fee_adjustment_review", why="per key", policy_ambiguous=False) for i in case.expected_policy_ids],
    )
    synced = sync(account, policy, derived, case.expected_category, sections)
    synthesis = SynthesisResult(draft=draft, route=draft.review_route) if draft else SynthesisResult(errors=["no draft"])
    filed = create_review_case(ds, case.account_id, case.expected_category, derived, synthesis.route)
    return Investigation(case.expected_category, account, derived, policy, synced, synthesis, filed)


def golden_echo(ds: Dataset, sections: dict[str, PolicySection], case: Scenario) -> Investigation:
    """The golden answer as an Investigation: the scorer's calibration, every axis must PASS."""
    cautious = case.expected_route in ("need_more_evidence", "manual_escalation")
    draft = Draft(
        review_route=case.expected_route, event_ids=case.expected_required_event_ids,
        policy_source_ids=case.expected_policy_ids, confidence="medium" if cautious else "high",
        open_questions=["Which event does the customer mean?"] if case.expected_route == "need_more_evidence" else [],
        case_note=case.rationale,
    )
    return investigation(ds, sections, case, key_findings(case), draft)


def empty_stub(ds: Dataset, sections: dict[str, PolicySection], case: Scenario) -> Investigation:
    """An agent that found nothing and drafted nothing: the scorer must fail it."""
    return investigation(ds, sections, case, None, None)


PRODUCERS = {"golden_echo": golden_echo, "empty_stub": empty_stub}


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="python -m dispute_agent.evaluation.golden", description=__doc__.splitlines()[0])
    p.add_argument("--data", default="data")
    mode = p.add_mutually_exclusive_group(required=True)
    mode.add_argument("--build", action="store_true", help=f"rewrite {GOLDEN_PATH}")
    mode.add_argument("--check", action="store_true", help="compare it to a rebuild from scenarios.xlsx")
    args = p.parse_args(argv)
    ds = load_dataset(args.data)
    if args.build:
        cases = build_cases(ds)
        write_golden(cases)
        print(f"wrote {GOLDEN_PATH}: {len(cases)} cases")
        return 0
    diffs = check(ds)
    print("\n".join(diffs) if diffs else f"{GOLDEN_PATH} matches scenarios.xlsx")
    return 1 if diffs else 0


if __name__ == "__main__":
    sys.exit(main())
